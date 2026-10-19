from pathlib import Path
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path.cwd()
snapshot=json.loads((root/"docs/testing/results/final-schema.json").read_text(encoding="utf-8-sig"))
order=["CAMPUS","DEPARTMENT","USER","CATEGORY","LOCATION","LOST_ITEM","FOUND_ITEM","MATCH","CLAIM","NOTIFICATION","MONGO_OUTBOX","SCHEMA_MIGRATION"]
tables=sorted(snapshot["tables"],key=lambda t:order.index(t["name"].upper()) if t["name"].upper() in order else len(order))
for t in tables:
    t["name"]=t["name"].upper()
    for k in t["keys"]:
        if k.get("referencedTable"):k["referencedTable"]=k["referencedTable"].upper()
def tidy(s): return re.sub(r"\s+"," ",str(s).replace("_utf8mb4","").replace("`","").replace(chr(92)+"'","'")).strip()
def keys(t,kind):
    grouped={}
    for k in t["keys"]:
        if k["type"]==kind and k["columnName"]:
            grouped.setdefault(k["name"],[]).append(k)
    return grouped
def notation(t):
    pk={k["columnName"] for group in keys(t,"PRIMARY KEY").values() for k in group}
    fk={k["columnName"]:k for group in keys(t,"FOREIGN KEY").values() for k in group}
    columns=[]
    for c in t["columns"]:
        label=c["name"]
        if c["name"] in pk: label+=" PK"
        if c["name"] in fk:
            k=fk[c["name"]];label+=" FK -> "+k["referencedTable"]+"."+k["referencedColumn"]
        if c["expression"]:label+=" GENERATED"
        columns.append(label)
    return t["name"]+"("+", ".join(columns)+")"
def constraint_lines(t):
    lines=[]
    for group in keys(t,"PRIMARY KEY").values():lines.append("PRIMARY KEY ("+", ".join(k["columnName"] for k in group)+")")
    for name,group in keys(t,"FOREIGN KEY").items():
        k=group[0];lines.append(name+": ("+", ".join(x["columnName"] for x in group)+") -> "+k["referencedTable"]+"("+", ".join(x["referencedColumn"] for x in group)+"); ON UPDATE "+k["updateRule"]+"; ON DELETE "+k["deleteRule"])
    for name,group in keys(t,"UNIQUE").items():lines.append(name+": UNIQUE ("+", ".join(k["columnName"] for k in group)+")")
    for c in t["checks"]:lines.append(c["name"]+": CHECK "+tidy(c["expression"]))
    for c in t["columns"]:
        if c["expression"]:lines.append(c["name"]+": STORED GENERATED AS "+tidy(c["expression"]))
    grouped={}
    for i in t["indexes"]:
        if i["nonUnique"]:grouped.setdefault(i["name"],[]).append(i["columnName"])
    for name,columns in grouped.items():lines.append(name+": INDEX ("+", ".join(columns)+")")
    return lines
notes={
"LOST_ITEM":"revision starts at 1 and the application increments it when editable core or MongoDB details change. The SQL date triggers reject a timestamp more than one minute in the future.",
"FOUND_ITEM":"revision starts at 1 and the application increments it when editable core or MongoDB details change. The SQL date triggers reject a timestamp more than one minute in the future.",
"MATCH":"lost_revision and found_revision snapshot report versions. sp_review_match checks them before acceptance. For canonical ACCEPTED rows, generated unique report IDs allow at most one accepted match per lost report and per found report. Pending/rejected rows have NULL generated values. Review-role triggers require a STAFF or ADMIN reviewer.",
"CLAIM":"For canonical APPROVED rows, approved_found_id permits at most one approved claim per found report. Other rows have NULL. Processor-role triggers require STAFF or ADMIN. The application/procedure also requires the accepted lost owner and prohibits processing one's own claim.",
"MONGO_OUTBOX":"This internal table stores ordered MongoDB UPSERT/DELETE payloads. document_id is not a SQL foreign key because its target depends on collection_name and a DELETE can outlive its domain row. Application services validate linked IDs.",
"SCHEMA_MIGRATION":"This internal table tracks applied migration names, checksums and timestamps. The current setup checks whether a migration name already exists; it does not compare an existing checksum or automatically reverse a partly applied DDL migration."
}
intro="The final implementation contains ten domain tables and two infrastructure tables. This schema is derived from a fresh database created by the submitted setup script, including all migrations. It includes exact columns, keys, checks, generated columns, defaults and indexes."
er_note="The conceptual ER model focuses on the core domain entities. Infrastructure tables and implementation-level concurrency/revision fields are documented separately in the final relational schema. Names use the uppercase spelling in the submitted SQL; Windows MySQL metadata stores the equivalent names in lowercase. A dash in Default means no non-NULL default is reported."
md=["# Final relational schema","",intro,"",er_note,"","PK means primary key. FK means foreign key. NULL indicates optional data. AUTO_INCREMENT keys are generated identities; STORED GENERATED columns are computed expressions. All listed domain foreign keys use RESTRICT for updates and deletes.","","Canonical fresh schema uses unsigned domain IDs. The compatibility upgrade preserves signed key types already present in the original local database while enabling missing auto-increment; the application does not change existing identities.",""]
dictionary=["# Final data dictionary","",intro,"",er_note,"","This dictionary uses the canonical fresh schema. Existing signed legacy identity types are retained by setup. Dates are persisted as UTC and displayed in IST. BOOLEAN is physically reported by MySQL as tinyint(1). Null/default information below is taken from information_schema.",""]
doc=Document()
sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.75)
for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Heading 3"]:
    style=doc.styles[name];style.font.name="Arial";style.font.color.rgb=RGBColor(0,0,0)
doc.styles["Normal"].font.size=Pt(11);doc.styles["Normal"].paragraph_format.space_after=Pt(7)
doc.styles["Title"].font.size=Pt(24)
for name in ["Heading 1","Heading 2","Heading 3"]:doc.styles[name].paragraph_format.keep_with_next=True
doc.add_paragraph("Campus Lost & Found Intelligent Matching System","Title")
doc.add_paragraph("Final Relational Schema","Subtitle")
doc.add_paragraph("DBMS Level 3 mini project | 2 October 2026")
doc.add_paragraph("Harshit Chandak  PES2UG24CS185\nHamza Shabir Sahapurwala  PES2UG24CS177")
doc.add_paragraph(intro);doc.add_paragraph(er_note)
doc.add_paragraph("PK denotes a primary key; FK denotes a foreign key. The source SQL and migrations are authoritative. The original local signed ID types are retained by the compatibility upgrade; the fresh schema below uses unsigned domain identities.")
p=sec.footer.paragraphs[0];p.alignment=2
run=p.add_run("Page ");run.font.size=Pt(9)
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");p._p.append(field)
current=None
for n,t in enumerate(tables,1):
    doc.add_page_break()
    if t["kind"]!=current:
        current=t["kind"];heading="A Domain Tables" if current=="domain" else "B Infrastructure and Internal Tables"
        doc.add_heading(heading,level=1);md.extend(["## "+heading,""]);dictionary.extend(["## "+heading,""])
    doc.add_heading(t["name"].replace("_"," "),level=2)
    p=doc.add_paragraph(notation(t));p.paragraph_format.space_after=Pt(10)
    for r in p.runs:r.font.name="Consolas";r.font.size=Pt(9)
    md.extend(["### "+t["name"],"","`"+notation(t)+"`","","| Column | Type | NULL | Default | Key or generation |","|---|---|---|---|---|"])
    dictionary.extend(["### "+t["name"],"","","| Column | Type | NULL | Default | Key or generation |","|---|---|---|---|---|"])
    table=doc.add_table(rows=1,cols=5);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    widths=[1.62,1.28,.5,1.15,2.45]
    for c,w in zip(table.columns,widths):c.width=Inches(w)
    headers=["Column","Type","NULL","Default","Key or generation"]
    for c,h,w in zip(table.rows[0].cells,headers,widths):c.text=h;c.width=Inches(w)
    pk={k["columnName"] for group in keys(t,"PRIMARY KEY").values() for k in group}
    fks={k["columnName"]:k for group in keys(t,"FOREIGN KEY").values() for k in group}
    for c in t["columns"]:
        flags=[]
        if c["name"] in pk:flags.append("PK")
        if c["name"] in fks:
            k=fks[c["name"]];flags.append("FK -> "+k["referencedTable"]+"."+k["referencedColumn"])
        if "auto_increment" in c["extra"]:flags.append("AUTO_INCREMENT")
        if c["expression"]:flags.append("STORED GENERATED")
        default="-" if c["defaultValue"] is None else str(c["defaultValue"])
        row=[c["name"],c["type"],c["nullable"],default,"; ".join(flags) or "-"]
        for cell,text,w in zip(table.add_row().cells,row,widths):cell.text=text;cell.width=Inches(w)
        escaped=[v.replace("|","&#124;") for v in row]
        line="| "+" | ".join(escaped)+" |";md.append(line);dictionary.append(line)
    for index,row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"DCE6F1" if index==0 else ("F5F7F9" if index%2==0 else "FFFFFF"));pr.append(shade)
            margins=OxmlElement("w:tcMar")
            for side in ["top","left","bottom","right"]:
                el=OxmlElement("w:"+side);el.set(qn("w:w"),"90");el.set(qn("w:type"),"dxa");margins.append(el)
            pr.append(margins)
            borders=OxmlElement("w:tcBorders")
            for side in ["top","left","bottom","right"]:
                b=OxmlElement("w:"+side);b.set(qn("w:val"),"single");b.set(qn("w:sz"),"4");b.set(qn("w:color"),"D9D9D9");borders.append(b)
            pr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.05
                for run in p.runs:run.font.size=Pt(9);run.bold=index==0
        no_split=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(no_split)
    repeat=OxmlElement("w:tblHeader");table.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph("Keys constraints and indexes","Heading 3");md.extend(["","Keys, constraints and indexes:",""]);dictionary.extend(["","Keys, constraints and indexes:",""])
    for line in constraint_lines(t):
        doc.add_paragraph(line,"List Bullet");md.append("- "+line);dictionary.append("- "+line)
    if t["name"].upper() in notes:
        doc.add_paragraph(notes[t["name"].upper()]);md.extend(["",notes[t["name"].upper()]]);dictionary.extend(["",notes[t["name"].upper()]])
    md.append("");dictionary.append("")
relationships=[
"CAMPUS 1 to many DEPARTMENT and LOCATION; each child campus_id is required.",
"DEPARTMENT 1 to many USER; USER.department_id is optional.",
"USER 1 to many LOST_ITEM, FOUND_ITEM, CLAIM submissions and NOTIFICATION recipients.",
"CATEGORY and LOCATION each classify or locate many lost/found reports; report references are required.",
"Each LOST_ITEM and FOUND_ITEM can have many candidate MATCH rows, but at most one canonical ACCEPTED match per report.",
"USER 1 to many MATCH reviews and CLAIM processing actions; actor fields are nullable before a decision and role-checked when supplied.",
"FOUND_ITEM 1 to many CLAIM rows; one per submitter/found pair and at most one canonical APPROVED claim per found item.",
"MONGO_OUTBOX and SCHEMA_MIGRATION have no SQL foreign keys. MongoDB links use the SQL IDs and are checked by the application, not cross-database FK constraints."
]
doc.add_page_break();doc.add_heading("Relationship cardinalities and implementation notes",1)
md.extend(["## Relationship cardinalities",""]);dictionary.extend(["## Relationships and design notes",""])
for line in relationships:doc.add_paragraph(line,"List Bullet");md.append("- "+line);dictionary.append("- "+line)
features="Verified inventory: 12 SQL base tables (10 domain and 2 internal), 2 views, 2 stored procedures, 1 stored function, 8 triggers and 4 MongoDB collections. Revisions are application-maintained fields; the database does not auto-increment them with a trigger. Terminal state rules are enforced through the APIs and stored procedures; raw SQL access must be restricted."
doc.add_paragraph(features);md.extend(["",features,"","Full installed metadata: [final-schema.json](../testing/results/final-schema.json). Executable sources: schema.sql, indexes_views.sql, advanced.sql, revisions.sql, procedures.sql and setup-databases.mjs.",""])
dictionary.extend(["",features,"","This data model separates reference entities from dependent reports and decisions. Generated uniqueness columns are implementation guards, not new conceptual entities. See [final relational schema](Relational_Schema.md), [Word schema](Relational_Schema.docx) and [fresh-schema metadata](../testing/results/final-schema.json). The original [ER diagram](ER_Diagram.png) and [editable ER](ER_Diagram.drawio) remain unchanged.",""])
doc.core_properties.title="Campus Lost & Found Intelligent Matching System Final Relational Schema"
doc.core_properties.author="Campus Lost and Found project"
doc.save(root/"docs/design/Relational_Schema.docx")
(root/"docs/design/Relational_Schema.md").write_text("\n".join(md),encoding="utf-8")
(root/"docs/design/Data_Dictionary.md").write_text("\n".join(dictionary),encoding="utf-8")
print("Updated DOCX and Markdown schema plus exact field-level dictionary for "+str(len(tables))+" tables.")
