"""Audit submitted text and document contents without printing credential values."""
from pathlib import Path
import argparse,json,re,zipfile,io,hashlib
from datetime import datetime,timezone
from urllib.parse import urlsplit,unquote
from xml.etree import ElementTree
from pypdf import PdfReader
ROOT=Path.cwd()
ROOT_FILES=[".env.example",".gitignore","README.md","package.json","package-lock.json","tsconfig.json","next-env.d.ts","next.config.mjs","vitest.config.ts","middleware.ts"]
DIRS=["app","components","lib","database","scripts","tests","docs"]
EXCLUDED_DIRS={".git","node_modules",".next",".next-dev","__pycache__","coverage"}
EXCLUDED_SCRIPTS={"read-requirements.py","prepare-lock.mjs","inspect-db.mjs"}
def selected_files():
    for name in ROOT_FILES:
        yield Path(name)
    for d in DIRS:
        for p in Path(d).rglob("*"):
            if not p.is_file() or any(x in EXCLUDED_DIRS for x in p.parts):continue
            if p.name.startswith(".env") or p.suffix in {".pyc",".tsbuildinfo",".tmp",".bak"}:continue
            if p.parts[0]=="scripts" and p.name in EXCLUDED_SCRIPTS:continue
            if p.suffix==".log" and not p.as_posix().startswith("docs/testing/results/"):continue
            yield p
def local_credentials():
    env={}
    for line in (ROOT/".env.local").read_text(encoding="utf-8-sig").splitlines():
        if not line.strip() or line.lstrip().startswith("#") or "=" not in line:continue
        k,v=line.split("=",1);env[k.strip()]=v.strip().strip("'\"")
    secrets=[]
    for key in ["MYSQL_PASSWORD","SESSION_SECRET","GOOGLE_CLIENT_SECRET","GOOGLE_CLIENT_ID"]:
        value=env.get(key,"")
        if value and not re.match(r"^(your_|replace_)",value):secrets.append((key,value))
    uri=env.get("MONGODB_URI","")
    if uri:
        password=urlsplit(uri).password
        if password:secrets.append(("MONGODB_URI_PASSWORD",unquote(password)))
    return secrets
def text_from(name,data):
    ext=Path(name).suffix.lower()
    if ext==".docx":
        texts=[]
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for part in z.namelist():
                if part.endswith(".xml"):
                    raw=z.read(part).decode("utf-8");texts.append(raw)
                    try:texts.append("".join(ElementTree.fromstring(raw).itertext()))
                    except ElementTree.ParseError:pass
        return "\n".join(texts)
    if ext==".pdf":return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(data)).pages)
    if ext in {".png",".jpg",".jpeg",".gif",".ico"}:return None
    try:return data.decode("utf-8-sig")
    except UnicodeDecodeError:return None
def audit(entries,secrets):
    findings=[];scanned=0;documents=0;envfiles=[]
    for name,data in entries:
        if Path(name).name.startswith(".env"):
            envfiles.append(name)
            if name!=".env.example":findings.append({"path":name,"reason":"real environment file"})
        text=text_from(name,data)
        if text is None:continue
        scanned+=1;documents+=Path(name).suffix.lower() in {".pdf",".docx"}
        for key,value in secrets:
            if len(value)>=8:
                hit=value in text or re.sub(r"\s+","",value) in re.sub(r"\s+","",text)
            else:
                hit=bool(re.search(r"(?:password|secret|MYSQL_PASSWORD)\s*[:=]\s*['\"]?"+re.escape(value)+r"(?:['\"\s,}]|$)",text,re.I))
            if hit:findings.append({"path":name,"reason":"configured credential value","key":key})
        if re.search(r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----",text):
            findings.append({"path":name,"reason":"private key block"})
        if re.search(r"\b(?:ghp_[A-Za-z0-9]{30,}|github_pat_[A-Za-z0-9_]{40,}|AIza[A-Za-z0-9_-]{30,})\b",text):
            findings.append({"path":name,"reason":"credential token pattern"})
    return {"status":"PASS" if not findings else "FAIL","textFilesScanned":scanned,"documentsExtracted":documents,"environmentFiles":envfiles,"findings":findings,"credentialValuesPrinted":False}
def validate_documents():
    from docx import Document
    snapshot=json.loads((ROOT/"docs/testing/results/final-schema.json").read_text(encoding="utf-8-sig"))
    expected_order=["CAMPUS","DEPARTMENT","USER","CATEGORY","LOCATION","LOST_ITEM","FOUND_ITEM","MATCH","CLAIM","NOTIFICATION","MONGO_OUTBOX","SCHEMA_MIGRATION"]
    metadata=sorted(snapshot["tables"],key=lambda t:expected_order.index(t["name"].upper()))
    word=Document(ROOT/"docs/design/Relational_Schema.docx")
    assert len(word.tables)==12,"Word schema table count differs"
    for table,meta in zip(word.tables,metadata):
        actual=[row.cells[0].text for row in table.rows[1:]]
        assert actual==[c["name"] for c in meta["columns"]],"Word column inventory differs"
    formatting_path=ROOT/"docs/testing/results/source-formatting.json"
    formatting=json.loads(formatting_path.read_text(encoding="utf-8-sig"))["files"] if formatting_path.exists() else []
    verified_formatting=[]
    for src in snapshot["sourceFiles"]:
        actual=hashlib.sha256((ROOT/src["path"]).read_bytes()).hexdigest()
        if actual==src["sha256"]:continue
        # Keep original capture hashes intact. Accept only a separately inspected
        # formatting version whose before/after byte hashes and full syntax-tree
        # comparison were recorded; any subsequent source edit fails this check.
        proof=next((p for p in formatting if p["path"]==src["path"]
            and p["originalSha256"]==src["sha256"] and p["currentSha256"]==actual
            and p.get("formattingOnly") is True and p.get("modelDefinitionsIdentical") is True
            and p.get("programStructureIdentical") is True
            and p["originalProgramStructureSha256"]==p["currentProgramStructureSha256"]),None)
        assert proof is not None,"Schema source changed since metadata capture"
        verified_formatting.append(src["path"])
    report=(ROOT/"docs/FINAL_REPORT.md").read_text(encoding="utf-8-sig")
    assert len(re.findall(r"^## \d+ ",report,re.M))==37,"Report section count differs"
    return {"wordColumnInventory":"PASS","schemaSourceHashes":"PASS","verifiedFormattingOnlySources":verified_formatting,"reportSections":37,"wordVisualLayout":"NOT RUN - dependency unavailable"}
def main():
    p=argparse.ArgumentParser();p.add_argument("--archive");p.add_argument("--output",default="docs/testing/results/secret-scan.json");args=p.parse_args()
    if args.archive:
        with zipfile.ZipFile(args.archive) as z:
            entries=[(n,z.read(n)) for n in z.namelist() if not n.endswith("/")]
    else:entries=[(p.as_posix(),p.read_bytes()) for p in selected_files()]
    result={"checkedAt":datetime.now(timezone.utc).isoformat(),"scope":"archive" if args.archive else "submission source and extracted documents",**audit(entries,local_credentials()),"documentConsistency":validate_documents()}
    dest=ROOT/args.output;dest.parent.mkdir(parents=True,exist_ok=True);dest.write_text(json.dumps(result,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"status":result["status"],"scope":result["scope"],"textFilesScanned":result["textFilesScanned"],"documentsExtracted":result["documentsExtracted"],"findingCount":len(result["findings"])}))
    raise SystemExit(0 if result["status"]=="PASS" else 1)
if __name__=="__main__":main()
